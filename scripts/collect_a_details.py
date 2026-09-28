# -*- coding: utf-8 -*-
"""
全A明细采集器 collect_a_details.py —— 云端替代「手动通达信导出 全部Ａ股{D}.xlsx」
================================================================================
背景：原 全部Ａ股{D}.xlsx（5227×23）为手动从通达信导出，是板块打分/竞价校准的地基，
不可自动化。本脚本用东财 push2 clist（与 fetch_snapshot.py 同源）重建同一份数据。

双时点设计（核心）：
  --mode auction  09:26 cron 运行：抓竞价成交额（届时 f6=竞价金额，**收盘后不可回补**）
                  → 存 全A竞价快照_{D}.json（原子写）
  --mode full     15:40 cron 运行：全字段快照 + 合并当日竞价快照 + 回读昨日快照
                  → 产出 全部Ａ股{D}.xlsx（同名同列序，下游 build_* 零改动）
                  ＋ 全A快照_{D}.json（自累积：供次日 昨成交额/昨涨幅%/昨竞价额）
  --mode compare --old <旧xlsx> --new <新xlsx>   P3 双跑验收：逐列对比报告

列口径（对照手动 TDX 导出 20260924 实测格式）：
  代码      6 位字符串（下游 str().strip() 兼容 ="..." 公式串）
  总金额/昨成交额/主力净额/开盘金额   **万元**（build_sector_full 注释确认金额单位万元）
  流通市值  文本「2192.84亿」格式（与手动导出一致）
  均价      总金额(元)/(成交量(手)×100)
  开盘%     (今开/昨收-1)×100 —— 与 TDX「集合竞价开盘涨幅」口径一致
  竞价昨比  今竞价金额/昨竞价金额（自家快照差分；⚠️ TDX 原口径未 100% 证实，
            但下游仅作透传标签（build_auction_baseline jjzb→next_jjzb），不进打分；
            1J2 的竞昨比门自行计算，不读此列）→ P3 双跑期用 TDX 值反证校准
  开盘换手Z/竞价涨停买/昨涨幅%   下游无脚本使用；尽力填充/占位
  细分行业  东财 f100 行业名（⚠️ 口径变更：TDX 350 细分行业 → 东财行业，
            sector_taxonomy 需配东财映射表，46 L1 归口逻辑不变）

环境变量：DETAILS_DATA_DIR（默认 D:\\OneDrive\\Stock\\details）
用法：
  python collect_a_details.py --mode auction --date 20260929
  python collect_a_details.py --mode full    --date 20260929
  python collect_a_details.py --mode compare --old a.xlsx --new b.xlsx
"""
from __future__ import annotations

import argparse
import datetime
import json
import os
import sys
import time

import requests

DETAILS_DIR = os.environ.get("DETAILS_DATA_DIR", r"D:\OneDrive\Stock\details")
EM_UT = "bd1d9ddb040334b8a88925d7c85c2f89"
SCAN_FS = "m:0+t:6,m:0+t:80,m:1+t:2,m:1+t:23"   # 沪深A股（不含北交所，与手动导出口径一致）
AUCTION_FIELDS = "f6,f8,f12,f14,f17,f18"
FULL_FIELDS = "f2,f3,f5,f6,f7,f8,f12,f14,f15,f16,f17,f18,f20,f21,f62,f100,f102"
HEADERS = ["代码", "名称", "细分行业", "现价", "涨幅%", "涨跌", "总金额", "昨成交额",
           "主力净额", "流通市值", "今开", "最高", "最低", "均价", "振幅%", "换手Z",
           "开盘%", "开盘金额", "开盘换手Z", "竞价昨比", "竞价涨停买", "昨涨幅%", "地区"]

session = requests.Session()
session.trust_env = False
session.headers.update({"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})


def get_with_retry(url, retries=3, backoff=2, **kw):
    for i in range(retries):
        try:
            return session.get(url, timeout=20, **kw)
        except requests.RequestException as e:
            if i == retries - 1:
                raise
            time.sleep(backoff * (i + 1))
    raise RuntimeError("unreachable")


def _f(v, default=None):
    try:
        x = float(v)
        return x if x != 0 or v == 0 else default
    except (TypeError, ValueError):
        return default


def scan_clist(fields):
    """全市场分页扫描 → [{code,name,...原始字段}]（fid=f12 升序，翻页稳定）"""
    out, pn = [], 1
    while pn <= 60:   # clist 按字段数限流：实测 pz=1000 也只返回 100 行/页 → 全市场 ≈56 页
        url = ("http://push2.eastmoney.com/api/qt/clist/get?pn=%d&pz=1000&po=1&np=1&fltt=2"
               "&invt=2&fid=f12&fs=%s&fields=%s&ut=%s" % (pn, SCAN_FS, fields, EM_UT))
        r = get_with_retry(url)
        data = (r.json().get("data") or {})
        diff = data.get("diff") or []
        if isinstance(diff, dict):
            diff = list(diff.values())
        if not diff:
            break
        out.extend(diff)
        total = int(data.get("total") or 0)
        print(f"  page {pn}: +{len(diff)} (total={total})")
        if len(out) >= total:
            break
        pn += 1
        time.sleep(0.3)
    return out


def atomic_json(payload, path):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)
    os.replace(tmp, path)


def load_json(path):
    if not path or not os.path.isfile(path):
        return None
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return None


def prev_store_path(date):
    """最近一个 < date 的 全A快照_{D}.json"""
    best = None
    try:
        names = os.listdir(DETAILS_DIR)
    except OSError:
        return None
    for n in names:
        if n.startswith("全A快照_") and n.endswith(".json"):
            d = n[5:-5]
            if len(d) == 8 and d.isdigit() and d < date and (best is None or d > best):
                best = d
    return os.path.join(DETAILS_DIR, f"全A快照_{best}.json") if best else None


# ---------------- auction 模式 ----------------
def run_auction(date):
    rows = scan_clist(AUCTION_FIELDS)
    if len(rows) < 4000:
        print(f"[FATAL] 竞价快照仅 {len(rows)} 行（<4000），疑似接口异常，拒绝落盘", file=sys.stderr)
        return 2
    payload = {
        "date": date,
        "collected_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "rows": {str(d["f12"]).zfill(6): {"jj_amt": _f(d.get("f6")), "hsl": _f(d.get("f8"))}
                 for d in rows},
    }
    path = os.path.join(DETAILS_DIR, f"全A竞价快照_{date}.json")
    atomic_json(payload, path)
    print(f"[auction] 已写 {path}（{len(payload['rows'])} 只）")
    return 0


# ---------------- full 模式 ----------------
def run_full(date, force=False):
    xlsx_path = os.path.join(DETAILS_DIR, f"全部Ａ股{date}.xlsx")
    if os.path.exists(xlsx_path) and not force:
        print(f"已存在 {xlsx_path}，跳过（--force 覆盖）")
        return 0

    auction = load_json(os.path.join(DETAILS_DIR, f"全A竞价快照_{date}.json")) or {}
    if not auction:
        print("[WARN] 当日无竞价快照 → 开盘金额/开盘换手Z/竞价昨比置空（勿静默填补）")
    arows = auction.get("rows", {})

    prev = load_json(prev_store_path(date)) or {}
    prows = prev.get("rows", {})
    if prev:
        print(f"[full] 昨日快照：{prev.get('date')}（{len(prows)} 只）")
    else:
        print("[WARN] 无昨日全A快照 → 昨成交额/昨涨幅% 置空（首日正常，次日起自动补齐）")

    rows = scan_clist(FULL_FIELDS)
    if len(rows) < 4000:
        print(f"[FATAL] 全量快照仅 {len(rows)} 行（<4000），疑似接口异常，拒绝落盘", file=sys.stderr)
        return 2

    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADERS)
    store = {}
    n = 0
    for d in rows:
        code = str(d.get("f12", "")).zfill(6)
        name = d.get("f14") or ""
        px, chg = _f(d.get("f2")), _f(d.get("f3"))
        prevc = _f(d.get("f18"))
        amt = _f(d.get("f6"))            # 元
        vol = _f(d.get("f5"))            # 手
        net = _f(d.get("f62"))           # 元
        cap = _f(d.get("f21"))           # 元
        hi, lo, op = _f(d.get("f15")), _f(d.get("f16")), _f(d.get("f17"))
        ind = d.get("f100") or "--"
        region = (d.get("f102") or "").replace("板块", "") or "--"
        a = arows.get(code, {})
        p = prows.get(code, {})
        jj_amt, jj_hsl = a.get("jj_amt"), a.get("hsl")
        prev_amt, prev_chg = p.get("f6"), p.get("f3")
        prev_jj = p.get("jj_amt")
        row = [
            code, name, ind, px, chg,
            round(px - prevc, 2) if (px is not None and prevc) else None,
            round(amt / 1e4, 2) if amt is not None else None,
            round(prev_amt / 1e4, 2) if prev_amt else None,
            round(net / 1e4, 2) if net is not None else None,
            f"{cap / 1e8:.2f}亿" if cap else "--",
            op, hi, lo,
            round(amt / (vol * 100), 2) if (amt and vol) else None,
            _f(d.get("f7")), _f(d.get("f8")),
            round((op / prevc - 1) * 100, 2) if (op and prevc) else None,
            round(jj_amt / 1e4, 2) if jj_amt else None,
            jj_hsl,
            (round(jj_amt / prev_jj, 2)
             if (jj_amt and prev_jj) else None),
            "--",
            prev_chg,
            region,
        ]
        ws.append(row)
        store[code] = {"f3": chg, "f6": amt, "f17": op, "f18": prevc, "jj_amt": jj_amt}
        n += 1
    tmp = xlsx_path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, xlsx_path)
    atomic_json({"date": date,
                 "collected_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                 "rows": store}, os.path.join(DETAILS_DIR, f"全A快照_{date}.json"))
    print(f"[full] 已写 {xlsx_path}（{n} 行）＋ 全A快照_{date}.json")
    return 0


# ---------------- compare 模式（P3 双跑验收） ----------------
def run_compare(old_path, new_path):
    import openpyxl

    def load(p):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        hdr = [str(x).strip() if x is not None else "" for x in next(it)]
        rows = {}
        for r in it:
            if not r or r[0] is None:
                continue
            code = str(r[0]).strip().zfill(6).lstrip("=").strip('"')
            rows[code] = dict(zip(hdr, r))
        wb.close()
        return rows

    old, new = load(old_path), load(new_path)
    common = set(old) & set(new)
    print(f"旧 {len(old)} 行 / 新 {len(new)} 行 / 交集 {len(common)} 只")
    only_old = set(old) - set(new)
    only_new = set(new) - set(old)
    if only_old:
        print(f"仅旧版有：{sorted(only_old)[:10]}{'...' if len(only_old) > 10 else ''}")
    if only_new:
        print(f"仅新版有：{sorted(only_new)[:10]}{'...' if len(only_new) > 10 else ''}")
    print(f"\n{'列':<8}{'可比值':>8}{'一致%':>8}{'中位偏差':>12}  说明")
    for col in HEADERS:
        if col in ("名称", "细分行业", "地区", "竞价涨停买"):
            diff_txt = sum(1 for c in common
                           if str(old[c].get(col) or "").strip() != str(new[c].get(col) or "").strip())
            print(f"{col:<8}{len(common):>8}{100 * (1 - diff_txt / max(len(common), 1)):>7.1f}%"
                  f"{'—':>12}  文本列，不一致 {diff_txt} 只")
            continue
        diffs, checked = [], 0
        for c in common:
            a, b = old[c].get(col), new[c].get(col)
            try:
                fa = float(str(a).replace("亿", "").replace(",", "").strip())
                fb = float(str(b).replace("亿", "").replace(",", "").strip())
            except (TypeError, ValueError):
                continue
            if abs(fa) < 1e-9 and abs(fb) < 1e-9:
                continue
            checked += 1
            diffs.append(abs(fa - fb))
        if not checked:
            print(f"{col:<8}{0:>8}{'—':>8}{'—':>12}  无可比数值（空列/占位）")
            continue
        diffs.sort()
        med = diffs[len(diffs) // 2]
        ok = sum(1 for x in diffs if x <= max(0.02, abs(med) * 0.05))
        print(f"{col:<8}{checked:>8}{100 * ok / checked:>7.1f}%{med:>12.4f}  中位绝对偏差")
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mode", choices=["auction", "full", "compare"], required=True)
    ap.add_argument("--date", default=datetime.date.today().strftime("%Y%m%d"))
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--old")
    ap.add_argument("--new")
    args = ap.parse_args()
    if args.mode == "compare":
        if not (args.old and args.new):
            print("compare 模式需要 --old 与 --new", file=sys.stderr)
            return 2
        return run_compare(args.old, args.new)
    os.makedirs(DETAILS_DIR, exist_ok=True)
    if args.mode == "auction":
        return run_auction(args.date)
    return run_full(args.date, force=args.force)


if __name__ == "__main__":
    raise SystemExit(main())
